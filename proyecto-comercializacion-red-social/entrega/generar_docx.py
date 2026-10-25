# -*- coding: utf-8 -*-
"""Genera documento APA 7 (trabajo estudiantil) — Geek Universe 3D."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
CAP = ROOT.parent / "capturas"
OUT = ROOT / "AmayaGil_Propuesta_Ecommerce.docx"

FONT = "Times New Roman"
BODY = Pt(12)
TABLE = Pt(11)


def _font(run, size=BODY, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def setup_document(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = BODY
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
        add_page_number_header(section)


def add_page_number_header(section):
    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    _font(run, size=Pt(12))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(0)
    p.alignment = align
    return p


def body(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = para(doc, align)
    r = p.add_run(text)
    _font(r)
    return p


def body_mixed(doc, parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """parts: list of (text, bold, italic)"""
    p = para(doc, align)
    for text, bold, italic in parts:
        r = p.add_run(text)
        _font(r, bold=bold, italic=italic)
    return p


def heading1(doc, text):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    _font(r, bold=True)
    return p


def heading2(doc, text):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    _font(r, bold=True)
    return p


def heading3(doc, text):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run(text)
    _font(r, bold=True, italic=True)
    return p


def table_label(doc, number, title):
    p = para(doc, WD_ALIGN_PARAGRAPH.LEFT)
    r1 = p.add_run(f"Tabla {number}\n")
    _font(r1, bold=True)
    r2 = p.add_run(title)
    _font(r2, italic=True)


def figure_label(doc, number, title):
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    r1 = p.add_run(f"Figura {number}\n")
    _font(r1, bold=True)
    r2 = p.add_run(title)
    _font(r2, italic=True)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            for run in p.runs:
                _font(run, size=TABLE, bold=True)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in p.runs:
                    _font(run, size=TABLE)
    body(doc, "")


def add_figure(doc, path: Path, number: int, title: str, note: str):
    figure_label(doc, number, title)
    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    if path.exists():
        run = p.add_run()
        run.add_picture(str(path), width=Inches(5.25))
    else:
        r = p.add_run("[Imagen no encontrada]")
        _font(r, italic=True)
    note_p = para(doc, WD_ALIGN_PARAGRAPH.LEFT)
    r = note_p.add_run(f"Nota. {note}")
    _font(r, size=Pt(11), italic=True)
    body(doc, "")


def reference_entry(doc, text_parts):
    """text_parts: list of (str, italic)"""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(0)
    for text, italic in text_parts:
        r = p.add_run(text)
        _font(r, italic=italic)


def title_page(doc):
    # Espacio superior (APA: título en mitad superior)
    for _ in range(4):
        para(doc, WD_ALIGN_PARAGRAPH.CENTER)

    p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
    title = (
        "Propuesta de Comercialización de un Producto "
        "por Medio de una Red Social"
    )
    r = p.add_run(title)
    _font(r, bold=True)

    for line in [
        "",
        "Alonso Esteban Amaya Gil",
        "",
        "Institución Universitaria Agustiniana",
        "",
        "Comercio electrónico (E-commerce): Unidades 1 y 2",
        "",
        "José Vicente Barreto Hernández",
        "",
        "11 de octubre de 2026",
    ]:
        p = para(doc, WD_ALIGN_PARAGRAPH.CENTER)
        if line:
            r = p.add_run(line)
            _font(r)

    doc.add_page_break()


def main():
    doc = Document()
    setup_document(doc)
    title_page(doc)

    heading1(doc, "Introducción")
    body(
        doc,
        "La comercialización por plataformas digitales permite a organizaciones y "
        "emprendimientos ampliar cobertura, reducir costos de intermediación y dialogar "
        "directamente con clientes (Kotler y Keller, 2016). En este contexto, las redes "
        "sociales dejaron de ser espacios únicamente informativos y pasaron a funcionar "
        "como canales de social commerce, en los que el contenido visual, las "
        "demostraciones del proceso productivo y la interacción con la audiencia "
        "influyen en la decisión de compra (Statista, 2024).",
    )
    body(
        doc,
        "El presente documento presenta una propuesta de comercio electrónico para "
        "comercializar productos de impresión 3D mediante Instagram, bajo el emprendimiento "
        "Geek Universe 3D (usuario @geekuniverse3d), perfil activo orientado a figuras de "
        "anime, superhéroes y videojuegos. Instagram concentra públicos interesados en "
        "cultura geek, coleccionismo y personalización, y ofrece herramientas para negocios "
        "(cuenta profesional, estadísticas y mensajes directos) (Instagram, 2024).",
    )
    body(
        doc,
        "El objetivo es describir la selección del canal, la configuración del perfil, la "
        "imagen de marca del emprendimiento y la presentación de un producto real de "
        "impresión 3D, relacionando estos elementos con los conceptos de las unidades 1 y 2: "
        "comercio electrónico, marketing digital y estrategia de canales en entornos digitales.",
    )

    heading1(doc, "Desarrollo de la Actividad")

    heading2(doc, "Selección de la Red Social")
    heading3(doc, "Análisis Comparativo")
    table_label(
        doc,
        1,
        "Comparación de Redes Sociales para la Comercialización de Figuras 3D",
    )
    add_table(
        doc,
        ["Criterio", "Instagram", "Facebook", "TikTok"],
        [
            [
                "Visualización del producto",
                "Alto (fotos, carruseles, Reels)",
                "Medio (Marketplace, grupos)",
                "Alto (video corto)",
            ],
            [
                "Público 15–40 años (coleccionismo)",
                "Muy presente",
                "Amplio, menor foco en nicho",
                "Creciente",
            ],
            [
                "Pedidos y cotizaciones",
                "Mensaje directo, enlace en biografía",
                "Messenger, Marketplace",
                "Mensaje directo y enlaces",
            ],
            [
                "Costo inicial",
                "Bajo (contenido orgánico)",
                "Bajo",
                "Bajo",
            ],
        ],
    )
    body(
        doc,
        "Decisión. Se selecciona Instagram como canal principal porque las figuras impresas "
        "se comercializan mejor cuando el cliente aprecia detalle, escala y acabado en "
        "imágenes de calidad. El perfil del emprendimiento ya existe en esta red, lo que "
        "permite documentar una propuesta auténtica (Laudon y Traver, 2021).",
    )

    heading3(doc, "Público Objetivo (Buyer Persona)")
    body(
        doc,
        "La buyer persona definida es Camila, de 24 años, fan de anime y videojuegos. "
        "Su motivación es coleccionar figuras de personajes icónicos (Pokémon, One Piece, "
        "superhéroes) con diseños diferenciados, sin depender de importaciones costosas. "
        "Su comportamiento digital incluye seguir cuentas de impresión 3D en Instagram y "
        "comprar tras evaluar fotografías de detalle y reseñas en comentarios o mensaje "
        "directo. La objeción principal es la calidad del acabado; el emprendimiento responde "
        "con fotografías de producto en alta definición, descripción del material impreso y "
        "atención personalizada.",
    )

    heading2(doc, "Creación del Perfil en la Red Social")
    heading3(doc, "Configuración del Perfil")
    body(
        doc,
        "El perfil de impresión 3D del emprendimiento ya fue creado en Instagram. La "
        "evidencia documentada incluye los siguientes elementos:",
    )
    items = [
        "Cuenta profesional con categoría Producto/servicio (ver Figura 1).",
        "Nombre comercial GeekUniverse3D y usuario @geekuniverse3d (5 publicaciones, "
        "20 seguidores y 82 cuentas seguidas al momento de la captura).",
        "Botones Seguir y Mensaje visibles, habilitando contacto comercial directo.",
        "Biografía orientada a tienda virtual de figuras de anime, superhéroes y videojuegos.",
    ]
    for i, item in enumerate(items, 1):
        p = para(doc)
        r = p.add_run(f"{i}.\t{item}")
        _font(r)

    body(doc, "URL del perfil: https://www.instagram.com/geekuniverse3d/")
    add_figure(
        doc,
        CAP / "perfil-completo.jpg",
        1,
        "Perfil de Instagram de Geek Universe 3D",
        "Captura de pantalla del perfil @geekuniverse3d, octubre de 2026. Fuente: elaboración propia.",
    )

    heading3(doc, "Justificación de la Decisión")
    body(
        doc,
        "Centralizar la vitrina en Instagram favorece coherencia de marca y continuidad "
        "entre lo publicado y lo descrito en la propuesta. La cuenta profesional permite "
        "consultar alcance e interacciones por publicación (Instagram, 2024).",
    )

    heading2(doc, "Elaboración de la Imagen de la Empresa en la Red Social")
    body(
        doc,
        "La identidad en red social debe comunicar personalización, calidad de impresión "
        "y un canal claro de pedidos (Kotler y Keller, 2016).",
    )
    heading3(doc, "Elementos de Marca Aplicados")
    table_label(doc, 2, "Identidad Digital de Geek Universe 3D en Instagram")
    add_table(
        doc,
        ["Elemento", "Aplicación en Instagram"],
        [
            ["Nombre comercial", "Geek Universe 3D"],
            [
                "Enfoque del negocio",
                "Figuras 3D de anime, superhéroes y videojuegos; coleccionables de nicho geek",
            ],
            [
                "Foto de perfil",
                "Logo circular con personajes de cultura pop y texto Geek Universe 3D",
            ],
            [
                "Estética del feed",
                "Grid con figuras impresas (Pokémon, One Piece, escenas de acción)",
            ],
            [
                "Biografía",
                "Bienvenida a GeekUniverse3D; figuras únicas; enlace a tienda virtual",
            ],
            [
                "Canal de contacto",
                "Botón Mensaje y referencia a tienda virtual en publicaciones",
            ],
        ],
    )
    body_mixed(
        doc,
        [
            ("Texto de biografía visible: ", False, False),
            (
                "«¡Bienvenido a GeekUniverse3D! Descubre figuras únicas de anime, "
                "superhéroes y videojuegos en nuestra tienda virtual…»",
                False,
                True,
            ),
        ],
    )

    heading3(doc, "Contenido Estratégico en el Perfil")
    body(
        doc,
        "El contenido publicado cumple tres funciones comerciales: (a) catálogo visual con "
        "figuras Pokémon, Tony Tony Chopper (One Piece) y escenas de acción; (b) producto "
        "estrella documentado (set Pokémon iniciales de Kanto y Pikachu); y (c) prueba de "
        "confianza mediante detalle fotográfico e interacción social. La publicación "
        "analizada incluye los hashtags #PokemonIniciales, #Pikachu, #3DPrints, "
        "#Collectibles y #GeekUniverse3D (Chaffey y Ellis-Chadwick, 2022).",
    )

    heading2(doc, "Presentación del Producto o Servicio")
    heading3(doc, "Descripción del Producto")
    body(
        doc,
        "Producto documentado: set de figuras 3D de Pokémon iniciales de Kanto y Pikachu "
        "(edición musculosa y coleccionable), impresas en material de color lila uniforme, "
        "presentadas sobre bases circulares.",
    )
    table_label(doc, 3, "Ficha del Producto Publicado en Instagram")
    add_table(
        doc,
        ["Atributo", "Detalle"],
        [
            ["Línea", "Coleccionables geek — anime y videojuegos"],
            ["Piezas", "Bulbasaur, Charmander, Squirtle y Pikachu"],
            ["Valor percibido", "Nostalgia Pokémon y pieza de exhibición"],
            ["Disponibilidad", "Tienda virtual del emprendimiento"],
            [
                "Canal de cierre",
                "Instagram (descubrimiento) y tienda virtual o mensaje directo (compra)",
            ],
            ["Fecha de publicación", "17 de enero de 2024"],
        ],
    )

    heading3(doc, "Publicación de Referencia")
    body_mixed(
        doc,
        [
            ("Extracto del texto publicado en @geekuniverse3d: ", False, False),
            (
                "«Captura la esencia de Kanto y desata la nostalgia con nuestras impresiones 3D "
                "de los Pokémon iniciales originales y Pikachu… en su versión musculosa… "
                "Disponibles ya en nuestra tienda virtual.»",
                False,
                True,
            ),
        ],
    )
    body(
        doc,
        "El llamado a la acción dirige a la tienda virtual y refuerza el posicionamiento "
        "mediante hashtags de coleccionismo (#3DPrints, #Collectibles).",
    )
    add_figure(
        doc,
        CAP / "publicacion-producto.jpg",
        2,
        "Publicación del Set Pokémon Iniciales Impresos en 3D",
        "Captura de la publicación en @geekuniverse3d, 17 de enero de 2024. Fuente: elaboración propia.",
    )

    heading3(doc, "Modelo de Comercialización")
    body(
        doc,
        "El emprendimiento opera bajo un modelo híbrido de social commerce y comercio "
        "electrónico ligero: Instagram funciona como vitrina y generador de tráfico emocional "
        "(nostalgia, fandom), mientras la tienda virtual concentra la conversión. Este esquema "
        "reduce costos de adquisición frente a tienda física y aprovecha la comunidad geek en "
        "red (Laudon y Traver, 2021).",
    )

    heading3(doc, "Indicadores de Seguimiento Propuestos")
    for item in [
        "Alcance e interacciones en publicaciones de producto.",
        "Número de mensajes directos clasificados como cotización o pedido confirmado.",
        "Tiempo promedio de respuesta en mensaje directo (meta: menos de 24 horas).",
    ]:
        p = para(doc)
        r = p.add_run(f"•\t{item}")
        _font(r)

    heading2(doc, "Síntesis de Decisiones")
    body(
        doc,
        "Se optó por Instagram por la naturaleza visual de las figuras impresas y por el "
        "perfil operativo Geek Universe 3D (geekuniverse3d). La imagen de marca combina logo "
        "geek, biografía de tienda virtual y grid temático; el producto ancla Pokémon "
        "iniciales Kanto más Pikachu demuestra la propuesta de valor; y el llamado a la "
        "acción enlaza la red social con la tienda virtual y el mensaje directo.",
    )

    heading1(doc, "Conclusiones")
    body(
        doc,
        "Comercializar figuras de impresión 3D mediante Instagram permitió a Geek Universe 3D "
        "posicionar un nicho claro (anime, superhéroes, videojuegos) y conectar emocionalmente "
        "con coleccionistas, como evidencia la publicación de Pokémon iniciales orientada a "
        "nostalgia y colección.",
    )
    body(
        doc,
        "La propuesta articula conceptos de comercio electrónico y marketing digital —canal, "
        "contenido, llamado a la acción, tienda virtual— con un caso real documentado mediante "
        "capturas y URL pública. Como mejora futura, convendría aumentar la frecuencia de "
        "publicaciones, explicitar precios en leyendas o tienda, y medir conversiones desde "
        "Instagram hacia la tienda virtual.",
    )
    body(
        doc,
        "En conjunto, la actividad cumple el objetivo formativo: demostrar una propuesta de "
        "comercialización digital coherente, apoyada en perfil real, imagen de marca geek y "
        "producto tangible publicado en la red social.",
    )

    heading1(doc, "Referencias")
    reference_entry(
        doc,
        [
            ("Chaffey, D., & Ellis-Chadwick, F. (2022). ", False),
            ("Digital marketing", True),
            (" (8th ed.). Pearson.", False),
        ],
    )
    reference_entry(
        doc,
        [
            ("Instagram. (2024). ", False),
            (
                "Tips and resources to help you grow your business on Instagram",
                True,
            ),
            (". Meta Platforms. https://business.instagram.com/", False),
        ],
    )
    reference_entry(
        doc,
        [
            ("Kotler, P., & Keller, K. L. (2016). ", False),
            ("Dirección de marketing", True),
            (" (15.ª ed.). Pearson Educación.", False),
        ],
    )
    reference_entry(
        doc,
        [
            ("Laudon, K. C., & Traver, C. G. (2021). ", False),
            ("E-commerce 2021: Business, technology, and society", True),
            (" (16th ed.). Pearson.", False),
        ],
    )
    reference_entry(
        doc,
        [
            ("Statista. (2024). ", False),
            ("Social commerce – Statistics & facts", True),
            (". https://www.statista.com/topics/8783/social-commerce/", False),
        ],
    )

    doc.save(str(OUT))
    print(f"Generado: {OUT}")


if __name__ == "__main__":
    main()
